"""Document handler for reading, redacting, and preserving formatting in DOCX and TXT files."""

import logging
from pathlib import Path
import time
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx import Document
from docx.text.paragraph import Paragraph

from src.config import OUTPUT_DIR
from src.pii_detector import PIIDetector
from src.redactor import RedactionResult, Redactor
from src.replacement_generator import ReplacementGenerator

logger = logging.getLogger(__name__)


class DocumentHandler:
    """Processes, redacts, and writes structured documents (DOCX, TXT) while preserving layout."""

    def __init__(
        self,
        detector: Optional[PIIDetector] = None,
        generator: Optional[ReplacementGenerator] = None,
    ) -> None:
        """Initialize document handler."""
        self.detector = detector or PIIDetector()
        self.generator = generator or ReplacementGenerator()
        self.redactor = Redactor(detector=self.detector, generator=self.generator)

    def _redact_paragraph(self, p: Paragraph) -> Tuple[int, Dict[str, int]]:
        """Redact PII in a single DOCX paragraph while preserving runs and formatting where possible.

        Args:
            p: The docx Paragraph object to redact in-place.

        Returns:
            Tuple of (detected_count, breakdown_dict).
        """
        full_text = p.text
        if not full_text.strip():
            return 0, {}

        entities = self.detector.detect(full_text)
        if not entities:
            return 0, {}

        breakdown: Dict[str, int] = {}
        for ent in entities:
            t = ent["type"]
            breakdown[t] = breakdown.get(t, 0) + 1

        # Fast path: single run or no runs
        if len(p.runs) <= 1:
            res = self.redactor.redact_text(full_text, entities=entities)
            if p.runs:
                p.runs[0].text = res.redacted_text
            else:
                p.text = res.redacted_text
            return len(entities), breakdown

        # Multi-run formatting preservation:
        # Map run character intervals (run_start, run_end)
        run_spans: List[Tuple[int, int]] = []
        cur_pos = 0
        for run in p.runs:
            rlen = len(run.text)
            run_spans.append((cur_pos, cur_pos + rlen))
            cur_pos += rlen

        # Check if any entity crosses across multiple runs
        crosses_runs = False
        for ent in entities:
            s, e = ent["start"], ent["end"]
            # Check which runs this entity intersects
            intersecting_runs = [
                i for i, (r_start, r_end) in enumerate(run_spans) if not (e <= r_start or s >= r_end)
            ]
            if len(intersecting_runs) > 1:
                crosses_runs = True
                break

        if not crosses_runs:
            # Each entity is fully contained within a single run: preserve all run styles perfectly!
            for i, (r_start, r_end) in enumerate(run_spans):
                run = p.runs[i]
                run_ents = [
                    {
                        "type": ent["type"],
                        "text": ent["text"],
                        "start": ent["start"] - r_start,
                        "end": ent["end"] - r_start,
                    }
                    for ent in entities
                    if r_start <= ent["start"] and ent["end"] <= r_end
                ]
                if run_ents:
                    # Redact run text
                    r_text = run.text
                    run_res = self.redactor.redact_text(r_text, entities=run_ents)
                    run.text = run_res.redacted_text
        else:
            # Entity spans cross run boundaries: fallback to paragraph-level redaction
            # while preserving paragraph style and alignment
            res = self.redactor.redact_text(full_text, entities=entities)
            p.text = res.redacted_text

        return len(entities), breakdown

    def process_docx(
        self,
        input_path: Path,
        output_path: Optional[Path] = None,
    ) -> Dict[str, Any]:
        """Redact a DOCX document in full (paragraphs, headings, lists, tables).

        Args:
            input_path: Path to the source DOCX file.
            output_path: Destination path for the redacted DOCX file.

        Returns:
            Dictionary containing processing metadata and statistics.
        """
        start_time = time.time()
        doc = Document(str(input_path))

        total_detected = 0
        combined_breakdown: Dict[str, int] = {}

        # 1. Process all main body paragraphs (headings, body text, lists)
        for p in doc.paragraphs:
            count, breakdown = self._redact_paragraph(p)
            total_detected += count
            for k, v in breakdown.items():
                combined_breakdown[k] = combined_breakdown.get(k, 0) + v

        # 2. Process all tables (cells and inner paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        count, breakdown = self._redact_paragraph(p)
                        total_detected += count
                        for k, v in breakdown.items():
                            combined_breakdown[k] = combined_breakdown.get(k, 0) + v

        # Determine output location
        if output_path is None:
            output_path = OUTPUT_DIR / f"redacted_{input_path.name}"

        doc.save(str(output_path))
        elapsed_sec = round(time.time() - start_time, 3)

        logger.info(
            "DOCX Redaction complete. Detected: %d PII items in %.3fs. Saved to: %s",
            total_detected,
            elapsed_sec,
            output_path,
        )

        return {
            "status": "success",
            "file_name": input_path.name,
            "output_path": str(output_path),
            "output_filename": output_path.name,
            "total_detected": total_detected,
            "breakdown": combined_breakdown,
            "processing_time_sec": elapsed_sec,
        }

    def process_txt(
        self,
        input_path: Path,
        output_path: Optional[Path] = None,
    ) -> Dict[str, Any]:
        """Redact a plain text file.

        Args:
            input_path: Path to the source text file.
            output_path: Destination path for the redacted text file.

        Returns:
            Dictionary containing processing metadata and statistics.
        """
        start_time = time.time()
        with open(input_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

        redaction_res = self.redactor.redact_text(content)

        if output_path is None:
            output_path = OUTPUT_DIR / f"redacted_{input_path.name}"

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(redaction_res.redacted_text)

        elapsed_sec = round(time.time() - start_time, 3)

        return {
            "status": "success",
            "file_name": input_path.name,
            "output_path": str(output_path),
            "output_filename": output_path.name,
            "total_detected": redaction_res.total_detected,
            "breakdown": redaction_res.breakdown,
            "processing_time_sec": elapsed_sec,
        }

    def process_document(
        self,
        input_path: Path,
        output_path: Optional[Path] = None,
    ) -> Dict[str, Any]:
        """Auto-detect format and redact document."""
        suffix = input_path.suffix.lower()
        if suffix == ".docx":
            return self.process_docx(input_path, output_path)
        elif suffix == ".txt":
            return self.process_txt(input_path, output_path)
        else:
            raise ValueError(f"Unsupported file format '{suffix}'. Supported formats: .docx, .txt")
