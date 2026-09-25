"""Automated test cases for DocumentHandler (.docx and .txt processing)."""

from pathlib import Path
import pytest
import docx

from src.config import SAMPLE_DIR, OUTPUT_DIR
from src.document_handler import DocumentHandler


class TestDocumentHandler:
    """Validate full document redaction workflows while preserving layout and structure."""

    def test_docx_redaction_workflow(self, doc_handler: DocumentHandler, tmp_path: Path):
        """Redact realistic sample ticket DOCX document and verify zero leakage."""
        sample_docx = SAMPLE_DIR / "sample_ticket_document.docx"
        output_docx = tmp_path / "test_redacted.docx"

        result = doc_handler.process_docx(sample_docx, output_docx)

        assert result["status"] == "success"
        assert result["total_detected"] > 0
        assert output_docx.exists()

        # Read back document and verify sensitive PII values do not exist
        doc = docx.Document(str(output_docx))
        doc_texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    doc_texts.append(cell.text)
        full_text = "\n".join(doc_texts)

        forbidden_pii = [
            "john.doe@gmail.com",
            "123-45-6789",
            "4111 1111 1111 1111",
            "rashmi.patil@example.com",
            "456-78-1234",
            "alex.mercer@cloudcorp.net",
            "219-45-7890",
        ]
        for pii in forbidden_pii:
            assert pii not in full_text, f"PII leak detected in DOCX: '{pii}'"

    def test_txt_redaction_workflow(self, doc_handler: DocumentHandler, tmp_path: Path):
        """Redact plain text document and verify zero leakage."""
        sample_txt = tmp_path / "sample.txt"
        sample_txt.write_text(
            "Ticket ID: TKT-9999\nCustomer: John Doe\nEmail: john.doe@gmail.com\nSSN: 123-45-6789\n",
            encoding="utf-8",
        )
        output_txt = tmp_path / "redacted_sample.txt"

        result = doc_handler.process_txt(sample_txt, output_txt)

        assert result["status"] == "success"
        assert result["total_detected"] >= 2
        assert output_txt.exists()

        redacted_content = output_txt.read_text(encoding="utf-8")
        assert "john.doe@gmail.com" not in redacted_content
        assert "123-45-6789" not in redacted_content
        assert "TKT-9999" in redacted_content  # Non-PII identifier preserved

    def test_unsupported_format_raises(self, doc_handler: DocumentHandler, tmp_path: Path):
        """Unsupported file extensions must raise ValueError."""
        invalid_file = tmp_path / "document.pdf"
        invalid_file.write_text("dummy", encoding="utf-8")

        with pytest.raises(ValueError) as exc:
            doc_handler.process_document(invalid_file)
        assert "Unsupported file format" in str(exc.value)
